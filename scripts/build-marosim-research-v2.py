"""Build the second research edition without rewriting previous document versions."""
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs/research"
sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location("research_layout", ROOT / "scripts/build-marosim-research-docs.py")
helpers = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helpers)


def add_table(document, rows):
    count = len(rows[0])
    is_characteristics = rows[0][0] == "важная характеристика"
    if is_characteristics:
        widths = [1.23, 1.65, 1.62, 2.0]
    elif count == 4 and rows[0][0].startswith("Категория"):
        widths = [1.48, 1.06, 0.74, 3.22]
    elif count == 4:
        widths = [1.12, 2.03, 1.25, 2.10]
    elif count == 5:
        widths = [1.0, 1.15, 1.50, 1.4, 1.45]
    elif count == 3:
        widths = [1.35, 2.55, 2.6]
    else:
        widths = [6.5 / count] * count
    table = document.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for i, (value, cell) in enumerate(zip(values, row.cells)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            helpers.set_cell_margins(cell, top=85, bottom=85, start=90, end=90)
            helpers.set_cell_shading(cell, "333333" if row_index == 0 else "F5F5F5" if row_index % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            helpers.add_inline(p, value, default_size=9.5, default_color="FFFFFF" if row_index == 0 else "000000")
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
            elif (count == 5 and i > 0) or (rows[0][0].startswith("Категория") and i in (1, 2)):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    helpers.set_repeat_table_header(table.rows[0])
    helpers.set_table_borders(table)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)


def build(source_name, output_name):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    helpers.configure_styles(doc)
    doc.styles["Title"].font.size = Pt(26)
    doc.styles["Title"].paragraph_format.space_before = Pt(0)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.05
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in list(doc.styles[name].element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Marosim  ·  ")
    helpers.set_run_font(run, 8, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = (DOCS / source_name).read_text(encoding="utf-8").splitlines()
    i = 0
    page_break_pending = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "<!-- PAGE -->":
            page_break_pending = True
            continue
        elif line.startswith("# "):
            p = doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("|"):
            rows = [helpers.split_table_row(line)]
            assert i < len(lines) and helpers.is_table_separator(lines[i])
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(helpers.split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
        else:
            p = doc.add_paragraph()
            if line.startswith("- "):
                p.style = "List Bullet"
                p.paragraph_format.space_after = Pt(5)
                line = line[2:]
            if line.startswith("**") and line.endswith("**") and line.count("**") == 2:
                p.paragraph_format.keep_with_next = True
            helpers.add_inline(p, line, default_size=11, default_color="000000")
        if page_break_pending:
            assert not line.startswith("|"), "Each new page must start with a heading or paragraph"
            p.paragraph_format.page_break_before = True
            page_break_pending = False
    doc.core_properties.author = "Marosim"
    doc.core_properties.title = lines[0][2:]
    doc.core_properties.subject = "Рынок, JTBD и проверка интерфейса Marosim"
    path = DOCS / output_name
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build("MAROSIM-MARKET-RESEARCH-UZBEKISTAN-v0.2.md", "Marosim_Market_Research_Uzbekistan_v0.2.docx")
    build("MAROSIM-USER-RESEARCH-QUESTIONNAIRE-UZ-CYR-v1.1.md", "Marosim_User_Research_Questionnaire_UZ_Cyr_v1.1.docx")
