from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(
    r"C:\Users\Lenovo\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-investment-committee-memo\assets\reference.docx"
)
OUTPUT_DIR = ROOT / "docs" / "research"
TEMP_DIR = ROOT / ".tmp" / "marosim-research-20260909" / "working"

BLACK = "111111"
DARK_GRAY = "5E5E5E"
MID_GRAY = "8A8A8A"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F5F5F5"
WHITE = "FFFFFF"


DOCUMENTS = [
    {
        "source": OUTPUT_DIR / "MAROSIM-USER-RESEARCH-QUESTIONNAIRE-UZ-CYR-v1.0.md",
        "output": OUTPUT_DIR / "Marosim_User_Research_Questionnaire_UZ_Cyr_v1.0.docx",
        "title": "Marosim фойдаланувчи\nинтервьюлари саволномаси",
        "subtitle": "Мижозлар ва таклиф муаллифлари учун ярим тузилган интервью ва маҳсулот тести",
        "header": "Marosim  |  Интервью саволномаси  |  09.09.2026",
        "meta": [
            ("Ким учун", "Тадқиқотчи ва суҳбат олиб борувчи", "Версия", "1.0"),
            ("Сегментлар", "Мижоз ва таклиф муаллифи", "Давомийлик", "40-50 дақиқа"),
            ("Тил", "Ўзбек тили кирилл ёзувида", "Маҳсулот", "Marosim"),
            ("Ҳолат", "Дала тадқиқоти учун тайёр", "Сана", "9 сентябрь 2026"),
        ],
        "subject": "Marosim билан мижозлар ва таклиф муаллифларини тадқиқ қилиш учун интервью саволномаси",
        "page_break_sections": set(range(2, 100)),
    },
    {
        "source": OUTPUT_DIR / "MAROSIM-MARKET-RESEARCH-UZBEKISTAN-v0.1.md",
        "output": OUTPUT_DIR / "Marosim_Market_Research_Uzbekistan_v0.1.docx",
        "title": "Исследование рынка\nи пользователей Marosim",
        "subtitle": "Узбекистан: свадебные и event-услуги, кабинетный этап и план полевой проверки",
        "header": "Marosim  |  Исследование рынка  |  09.09.2026",
        "meta": [
            ("Для кого", "Владелец продукта и команда", "Версия", "0.1"),
            ("Рынок", "Узбекистан", "Фокус", "Свадьбы и события"),
            ("Метод", "Кабинетный ресёрч и полевой план", "Статус", "Первый вывод готов"),
            ("Следующий шаг", "24 интервью и продуктовые тесты", "Дата", "9 сентября 2026"),
        ],
        "subject": "Первичный анализ рынка Marosim и план исследования двух сторон маркетплейса",
        "page_break_sections": set(range(2, 9)) | {10, 12},
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, size: float, *, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_separate, value, fld_end])
    set_run_font(run, 9, color=MID_GRAY)


def configure_page_chrome(document: Document, header_text: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    header = section.header
    for paragraph in header.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = header.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(header_text)
    set_run_font(run, 9, color=MID_GRAY)

    footer = section.footer
    for paragraph in footer.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = footer.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run("Marosim  |  Исследование продукта")
    set_run_font(run, 9, color=MID_GRAY)
    paragraph.add_run("\t")
    add_page_field(paragraph)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = document.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True

    for style_name, size, before, after in (("Heading 1", 18, 18, 7), ("Heading 2", 13, 13, 5)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Research Question" not in document.styles:
        question = document.styles.add_style("Research Question", WD_STYLE_TYPE.PARAGRAPH)
    else:
        question = document.styles["Research Question"]
    question.base_style = normal
    question.font.name = "Arial"
    question._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    question._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    question._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    question.font.size = Pt(11)
    question.font.bold = True
    question.font.color.rgb = RGBColor.from_string(BLACK)
    question.paragraph_format.space_before = Pt(8)
    question.paragraph_format.space_after = Pt(3)
    question.paragraph_format.keep_with_next = True

    for name in ("Research Bullet", "Research Number"):
        if name not in document.styles:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = document.styles[name]
        style.base_style = normal
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_after = Pt(3)


def add_hyperlink(paragraph, text: str, url: str, *, bold=False, italic=False, color=DARK_GRAY):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial")
    run_properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    run_properties.append(size)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    run_properties.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    if italic:
        run_properties.append(OxmlElement("w:i"))
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_PATTERN = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str, *, default_size=10.5, default_color=BLACK, default_italic=False) -> None:
    position = 0
    for match in TOKEN_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, default_size, color=default_color, italic=default_italic)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, default_size, bold=True, color=default_color, italic=default_italic)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url, color=DARK_GRAY)
        else:
            url = token.rstrip(".,;)")
            suffix = token[len(url) :]
            add_hyperlink(paragraph, url, url, color=DARK_GRAY)
            if suffix:
                run = paragraph.add_run(suffix)
                set_run_font(run, default_size, color=default_color, italic=default_italic)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, default_size, color=default_color, italic=default_italic)


def add_cover_bar(document: Document) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(4.95)
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = table.columns[index].width
        clear_paragraph(cell.paragraphs[0])
        set_cell_shading(cell, BLACK if index == 0 else WHITE)
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    row_height = OxmlElement("w:trHeight")
    row_height.set(qn("w:val"), "120")
    row_height.set(qn("w:hRule"), "exact")
    table.rows[0]._tr.get_or_add_trPr().append(row_height)
    set_table_borders(table, color=WHITE, size="0")


def add_cover(document: Document, config: dict) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(82)
    add_cover_bar(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lines = config["title"].split("\n")
    for index, line in enumerate(lines):
        if index:
            title.add_run().add_break()
        run = title.add_run(line)
        set_run_font(run, 34, bold=True, color=BLACK)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    add_inline(subtitle, config["subtitle"], default_size=14, default_color=DARK_GRAY)

    table = document.add_table(rows=len(config["meta"]), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.95, 2.35, 0.9, 2.3]
    for row_index, row_data in enumerate(config["meta"]):
        for col_index, value in enumerate(row_data):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
            set_cell_shading(cell, PALE_GRAY if col_index in (0, 2) else WHITE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, 9.5, bold=col_index in (0, 2), color=BLACK)
    set_table_borders(table)
    document.add_page_break()


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    width_sets = {
        2: [2.0, 4.5],
        3: [1.8, 2.35, 2.35],
        4: [1.35, 1.75, 1.55, 1.85],
        5: [1.15, 1.2, 0.75, 1.55, 1.85],
    }
    widths = width_sets.get(column_count, [6.5 / column_count] * column_count)
    for row_index, values in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if column_count >= 5:
                set_cell_margins(cell, top=70, start=70, bottom=70, end=70)
            else:
                set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, BLACK)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            else:
                set_cell_shading(cell, WHITE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 or column_count <= 3 else WD_ALIGN_PARAGRAPH.LEFT
            value = values[col_index] if col_index < len(values) else ""
            add_inline(
                paragraph,
                value,
                default_size=8.5 if column_count >= 5 else 9 if column_count >= 4 else 9.5,
                default_color=WHITE if row_index == 0 else BLACK,
            )
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.bold = True
    set_repeat_table_header(table.rows[0])
    set_table_borders(table)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def markdown_body_lines(path: Path) -> list[str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    title_seen = False
    subtitle_removed = False
    output = []
    for line in lines:
        if not title_seen and line.startswith("# "):
            title_seen = True
            continue
        if title_seen and not subtitle_removed and line.strip() and not line.startswith("#"):
            subtitle_removed = True
            continue
        output.append(line)
    return output


def add_markdown_body(document: Document, path: Path, page_break_sections: set[int]) -> None:
    lines = markdown_body_lines(path)
    section_number = 0
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        line = raw.strip()
        if not line:
            index += 1
            continue

        if line.startswith("## "):
            section_number += 1
            paragraph = document.add_paragraph(style="Heading 1")
            if section_number in page_break_sections:
                paragraph.paragraph_format.page_break_before = True
            paragraph.add_run(f"{section_number}. {line[3:].strip()}")
            for run in paragraph.runs:
                set_run_font(run, 18, bold=True)
            index += 1
            continue

        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.add_run(line[4:].strip())
            for run in paragraph.runs:
                set_run_font(run, 13, bold=True)
            index += 1
            continue

        if line.startswith("#### "):
            paragraph = document.add_paragraph(style="Research Question")
            add_inline(paragraph, line[5:].strip(), default_size=11)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_markdown_table(document, rows)
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="Research Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            marker = paragraph.add_run("• ")
            set_run_font(marker, 10.5, bold=True)
            add_inline(paragraph, line[2:].strip())
            index += 1
            continue

        if re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="Research Number")
            list_match = re.match(r"^(\d+)\.\s+(.*)$", line)
            marker = paragraph.add_run(f"{list_match.group(1)}. ")
            set_run_font(marker, 10.5, bold=True)
            add_inline(paragraph, list_match.group(2))
            index += 1
            continue

        if line.startswith("**") and line.endswith("**") and len(line) > 4:
            paragraph = document.add_paragraph(style="Research Question")
            add_inline(paragraph, line[2:-2], default_size=11)
            index += 1
            continue

        paragraph = document.add_paragraph(style="Normal")
        if line in {"Жавоб учун жой", "Қўшимча саволлар:"} or line.startswith("Қўшимча саволлар:") or line.startswith("Кузатинг:"):
            paragraph.paragraph_format.left_indent = Inches(0.18)
            add_inline(paragraph, line, default_size=9.5, default_color=DARK_GRAY, default_italic=True)
        else:
            add_inline(paragraph, line)
        index += 1


def build_document(config: dict) -> None:
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    working = TEMP_DIR / config["output"].name
    shutil.copyfile(REFERENCE, working)
    document = Document(working)
    clear_body(document)
    configure_styles(document)
    configure_page_chrome(document, config["header"])
    add_cover(document, config)
    add_markdown_body(document, config["source"], config["page_break_sections"])
    document.core_properties.title = config["title"].replace("\n", " ")
    document.core_properties.subject = config["subject"]
    document.core_properties.author = "Marosim"
    document.core_properties.keywords = "Marosim, market research, user research, Uzbekistan"
    document.save(config["output"])


def main() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    for config in DOCUMENTS:
        if not config["source"].exists():
            raise FileNotFoundError(config["source"])
        build_document(config)
        print(config["output"])


if __name__ == "__main__":
    main()
